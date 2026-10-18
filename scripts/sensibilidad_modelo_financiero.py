#!/usr/bin/env python3
"""Sensibilidad del escenario neutral: una variable por vez.

Para cada variable mueve su multiplicador de la hoja Supuestos (sens_*),
recalcula la planilla con LibreOffice y lee el VAN del escenario neutral. Informa
el VAN con la variable un 20 % abajo y un 20 % arriba, y el valor de equilibrio
(el que deja el VAN en cero), hallado por bisección. La tasa de descuento no
necesita recálculo: su valor de equilibrio es la TIR.

Escribe documento/chapters/tables/financiero-sensibilidad.tex. Tarda un par de
minutos, porque cada punto es un recálculo completo de la planilla.

Uso:  python scripts/sensibilidad_modelo_financiero.py   (con openpyxl y soffice)
"""
import os, subprocess, sys, tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import validar_modelo_financiero as v  # noqa: E402

PLANILLA = v.PLANILLA
SALIDA = os.path.join(v.SALIDA, 'financiero-sensibilidad.tex')
ITERACIONES = 10

# (multiplicador, variable, clave del supuesto que se informa, unidad, decimales)
VARIABLES = [
    ('sens_costo_analisis', 'Costo por análisis', 'costo_analisis', 'USD', 4),
    ('sens_precios', 'Precios B2B (nivel superior de la API)', 'precio_medios_grandes', 'USD/mes', 0),
    ('sens_captacion', 'Tasas de captación (medios de gran porte)', 'capt_medios_grandes', '%', 1),
]


def fila_de(ws, clave):
    return next(c.row for c in ws['A'] if c.value == clave)


def neutral(ws, clave):
    """Valor neutral de un supuesto: columna F si varía por escenario, D si es común."""
    r = fila_de(ws, clave)
    return ws[f'F{r}'].value if ws[f'D{r}'].value is None else ws[f'D{r}'].value


class Planilla:
    def __init__(self, directorio):
        self.dir = directorio
        self.wb = openpyxl.load_workbook(PLANILLA)
        self.sup = self.wb['Supuestos']

    def van(self, clave, multiplicador):
        celda = f'D{fila_de(self.sup, clave)}'
        self.sup[celda] = multiplicador
        entrada = os.path.join(self.dir, 'entrada', 'm.xlsx')
        os.makedirs(os.path.dirname(entrada), exist_ok=True)
        self.wb.save(entrada)
        self.sup[celda] = 1
        subprocess.run(['soffice', f'-env:UserInstallation=file://{self.dir}/lo', '--headless',
                        '--convert-to', 'xlsx', '--outdir', self.dir, entrada],
                       check=True, capture_output=True)
        hoja = openpyxl.load_workbook(os.path.join(self.dir, 'm.xlsx'), data_only=True)['Neutral']
        return next(f[1] for f in hoja.iter_rows(values_only=True) if f[0] == 'VAN (USD)')


def equilibrio(f, base_van):
    """Multiplicador con VAN = 0, buscado en tramos cada vez más alejados de 1."""
    tramos = [(0.8, 1.0), (1.0, 1.2), (0.5, 0.8), (1.2, 1.5), (1.5, 2.0), (0.2, 0.5), (2.0, 3.0)]
    cache = {1.0: base_van, 0.8: f(0.8), 1.2: f(1.2)}
    for lo, hi in tramos:
        for x in (lo, hi):
            if x not in cache:
                cache[x] = f(x)
        if cache[lo] * cache[hi] <= 0:
            f_lo = cache[lo]
            for _ in range(ITERACIONES):
                medio = (lo + hi) / 2
                f_medio = f(medio)
                if f_lo * f_medio <= 0:
                    hi = medio
                else:
                    lo, f_lo = medio, f_medio
            return (lo + hi) / 2, cache
    return None, cache


def variacion(x):
    return ('+' if x > 0 else '') + v.en_texto(100 * x, 1) + '\\,\\%'


def main():
    tasa, escenarios = v.leer(PLANILLA)
    base = escenarios['Neutral']
    base_van = base['VAN (USD)'][0]
    sup = openpyxl.load_workbook(PLANILLA, data_only=True)['Supuestos']
    filas = []
    with tempfile.TemporaryDirectory() as d:
        planilla = Planilla(d)
        for sens, nombre, clave, unidad, dec in VARIABLES:
            m, cache = equilibrio(lambda x: planilla.van(sens, x), base_van)
            valor = neutral(sup, clave)
            fmt = (lambda x: v.num(100 * x, dec) + '\\,\\%') if unidad == '%' else (
                lambda x: f'{v.num(x, dec)} {unidad}')
            eq = 'fuera del rango' if m is None else f'{fmt(valor * m)} ({variacion(m - 1)})'
            filas.append([nombre, fmt(valor), v.num(cache[0.8], 0), v.num(cache[1.2], 0), eq])
            print(f'  {nombre}: equilibrio ×{m}')
    neto = base['Flujo neto']
    tir = v.tir(neto)
    filas.append(['Tasa de descuento', v.num(100 * tasa, 0) + '\\,\\%',
                  v.num(v.van(neto, 0.8 * tasa), 0), v.num(v.van(neto, 1.2 * tasa), 0),
                  f'{v.num(100 * tir, 2)}\\,\\% ({variacion(tir / tasa - 1)})'])
    tex = v.tabla('sensibilidad', None,
                  'Sensibilidad del VAN del escenario neutral, en USD, moviendo una variable por vez. '
                  f'El VAN del caso base es de {v.en_texto(base_van)} USD. El valor de equilibrio es el '
                  'que anula el VAN y, entre paréntesis, su variación respecto del caso base',
                  'p{3.6cm} r r r p{3.4cm}',
                  ['Variable', 'Caso base', 'VAN $-$20\\,\\%', 'VAN +20\\,\\%',
                   'Equilibrio'], filas, letra='\\footnotesize')
    with open(SALIDA, 'w') as fh:
        fh.write(tex.replace('validar_modelo_financiero.py', 'sensibilidad_modelo_financiero.py'))
    print(f'  escrito  {os.path.relpath(SALIDA, v.ROOT)}')


if __name__ == '__main__':
    main()
