#!/usr/bin/env python3
"""Validación y exportación del modelo financiero del PFI.

Lee wiki/negocio/modelo-financiero.xlsx con los valores ya recalculados
(ver scripts/generar_modelo_financiero.py), recalcula por su cuenta los
indicadores de cada hoja de escenario y los compara con los de la planilla.
Si alguno no coincide, termina con código 1 e indica el indicador y el
escenario. Si todo coincide, escribe las tablas LaTeX de cada escenario en
documento/chapters/tables/, la del mercado, la de los supuestos (anexo) y
las cifras sueltas que cita el texto (financiero-cifras.tex, una macro por
cifra), para que ningún capítulo transcriba un número a mano.

Antes de validar corre un autochequeo con flujos de juguete de resultado
conocido, para que un error del cálculo de referencia no apruebe una
planilla rota.

Uso:  .venv/bin/python scripts/validar_modelo_financiero.py [--autochequeo]
          [--planilla RUTA] [--salida DIRECTORIO]
"""
import argparse, math, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANILLA = os.path.join(ROOT, 'wiki/negocio/modelo-financiero.xlsx')
SALIDA = os.path.join(ROOT, 'documento/chapters/tables')

NO_TIR = 'no definida'
NO_RECUPERA = 'no se recupera'


# ------------------------------------------------------------- indicadores

def descontar(flujos, tasa):
    return [f / (1 + tasa) ** t for t, f in enumerate(flujos)]


def van(flujos, tasa):
    return sum(descontar(flujos, tasa))


def tir(flujos, lo=-0.99, hi=100.0):
    """Tasa que anula el VAN, por bisección. Sin cambio de signo no hay TIR."""
    if min(flujos) >= 0 or max(flujos) <= 0:
        return None
    f_lo = van(flujos, lo)
    if f_lo * van(flujos, hi) > 0:
        return None
    for _ in range(200):
        medio = (lo + hi) / 2
        f_medio = van(flujos, medio)
        if f_lo * f_medio <= 0:
            hi = medio
        else:
            lo, f_lo = medio, f_medio
    return (lo + hi) / 2


def payback(flujos):
    """a + b/c: a = año anterior al recupero, b = saldo sin recuperar al
    cierre del año a, c = flujo del año del recupero."""
    acumulado = 0.0
    for t, f in enumerate(flujos):
        previo, acumulado = acumulado, acumulado + f
        if t > 0 and previo < 0 <= acumulado:
            return (t - 1) + (-previo) / f
    return NO_RECUPERA


def punto_equilibrio(costos, ingresos, clientes):
    """Clientes B2B necesarios para cubrir los costos del año, al ingreso medio."""
    if not ingresos or not clientes:
        return None
    return math.ceil(costos / (ingresos / clientes) - 1e-9)


def autochequeo():
    # Flujo A al 10 %: VAN = -100 + 54,5455 + 49,5868 = 4,1322.
    # TIR: 60x² + 60x - 100 = 0 => x = (-60 + √27.600)/120 = 0,884437 => 13,0662 %.
    # Acumulado -100, -40, 20 => 1 + 40/60. Descontado -100, -45,4545, 4,1322
    # => 1 + 45,4545/49,5868 = 1 + 60,5/66 = 1,916667.
    a = [-100, 60, 60]
    assert abs(van(a, 0.10) - 4.132231) < 1e-5
    assert abs(tir(a) - 0.130662) < 1e-5
    assert abs(payback(a) - 1.666667) < 1e-5
    assert abs(payback(descontar(a, 0.10)) - 1.916667) < 1e-5
    # Flujo B al 10 %: VAN = -100 + 30 × 2,486852 = -25,3944; no se recupera.
    b = [-100, 30, 30, 30]
    assert abs(van(b, 0.10) + 25.394440) < 1e-5
    assert payback(b) == NO_RECUPERA
    # Flujo C: nunca cambia de signo, no hay TIR.
    assert tir([-100, -10, -5]) is None
    # Flujo D: 110 / 1,10 = 100 => TIR = 10 % y VAN al 10 % = 0.
    assert abs(tir([-100, 110]) - 0.10) < 1e-9
    assert abs(van([-100, 110], 0.10)) < 1e-9
    # Punto de equilibrio: 1.000 / (300 / 4) = 13,33 => 14 clientes.
    assert punto_equilibrio(1000, 300, 4) == 14
    assert punto_equilibrio(1000, 0, 0) is None
    print('autochequeo: ok')


# ---------------------------------------------------------------- lectura

HOJAS_NO_ESCENARIO = {'Supuestos', 'Resumen'}


def leer(planilla):
    import openpyxl
    wb = openpyxl.load_workbook(planilla, data_only=True)
    sup = {f[0]: f[3] for f in wb['Supuestos'].iter_rows(values_only=True) if f[0]}
    escenarios = {}
    for hoja in wb.worksheets:
        if hoja.title in HOJAS_NO_ESCENARIO:
            continue
        filas = {f[0]: list(f[1:7]) for f in hoja.iter_rows(values_only=True) if f[0]}
        escenarios[hoja.title] = filas
    return sup['tasa_descuento'], escenarios


def leer_hoja(planilla, nombre):
    import openpyxl
    wb = openpyxl.load_workbook(planilla, data_only=True)
    return list(wb[nombre].iter_rows(values_only=True)) if nombre in wb.sheetnames else []


def van_con_residual(filas, tasa, g):
    """VAN más el valor residual de Gordon sobre el flujo operativo del año 5."""
    operativo = filas['Flujo neto'][-1] - (filas['Recupero del capital de trabajo'][-1] or 0)
    residual = max(0.0, operativo) * (1 + g) / (tasa - g)
    return van(filas['Flujo neto'], tasa) + residual / (1 + tasa) ** 5


def recalcular(filas, tasa, g=None):
    neto = filas['Flujo neto']
    extra = {}
    if 'VAN con valor residual (USD)' in filas and g is not None:
        extra['VAN con valor residual (USD)'] = van_con_residual(filas, tasa, g)
    return extra | {
        'VAN (USD)': van(neto, tasa),
        'TIR': tir(neto),
        'Payback simple (años)': payback(neto),
        'Payback descontado (años)': payback(descontar(neto, tasa)),
        'Punto de equilibrio (clientes B2B)': punto_equilibrio(
            filas['Costos totales'][-1], filas['Ingresos totales'][-1],
            filas['Clientes B2B (promedio del año, con contratos)'][-1]),
    }


TEXTO_NULO = {'TIR': NO_TIR, 'Punto de equilibrio (clientes B2B)': 'no definido'}


def coincide(planilla, propio):
    if isinstance(propio, str) or isinstance(planilla, str):
        return planilla == propio
    return abs(planilla - propio) <= 1e-6 * max(1, abs(propio)) + 1e-4


# ----------------------------------------------------------------- salida

CABECERA = '% Generado por scripts/validar_modelo_financiero.py — no editar a mano.\n'
FUENTE = 'Fuente: elaboración propia sobre la planilla del modelo financiero.'
SEGMENTOS = [  # (etiqueta en la planilla, encabezado de la tabla)
    ('Medios de gran porte', 'Medios de gran porte'),
    ('Medios socios de ADEPA (resto)', 'Resto de medios de ADEPA'),
    ('Verificadores', 'Verifica\\-dores'),
    ('Universidades y observatorios', 'Univer\\-sidades'),
    ('Agencias', 'Agencias'),
]


def num(v, dec=2):
    """Coma decimal, punto de miles y negativos entre paréntesis."""
    if abs(v) < 0.5 * 10 ** -dec:
        v = 0
    s = f'{abs(v):,.{dec}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return f'({s})' if v < 0 else s


def entero(v):
    return int(math.floor((v or 0) + 0.5))


def tabla(slug, clave, titulo, columnas, encabezado, filas, letra='\\small'):
    cuerpo = '\n'.join('    ' + ' & '.join(f) + r' \\' for f in filas)
    return CABECERA + f"""\\begin{{table}}[t]
  \\centering
  {letra}
  \\caption{{{titulo}. {FUENTE}}}
  \\label{{tab:financiero-{slug}{'-' + clave if clave else ''}}}
  \\begin{{tabular}}{{{columnas}}}
    \\hline
    {' & '.join(f'\\textbf{{{e}}}' for e in encabezado)} \\\\
    \\hline
{cuerpo}
    \\hline
  \\end{{tabular}}
\\end{{table}}
"""


def necesidad_maxima(filas):
    """Máximo saldo negativo del flujo neto acumulado y el año en que ocurre."""
    acumulado = filas['Flujo neto acumulado']
    anio = min(range(len(acumulado)), key=lambda t: acumulado[t])
    return max(0.0, -acumulado[anio]), anio


def filas_indicadores(filas):
    def valor(etiqueta, formato):
        v = filas[etiqueta][0]
        return v if isinstance(v, str) else formato(v)

    necesidad, anio = necesidad_maxima(filas)
    residual = ([['VAN con valor residual (USD)', valor('VAN con valor residual (USD)', num)]]
                if 'VAN con valor residual (USD)' in filas else [])
    return [
        ['VAN (USD)', valor('VAN (USD)', num)],
        *residual,
        ['TIR (\\%)', valor('TIR', lambda v: num(100 * v))],
        ['\\textit{Payback} (años)', valor('Payback simple (años)', num)],
        ['\\textit{Payback} descontado (años)', valor('Payback descontado (años)', num)],
        ['Punto de equilibrio en el año 5 (clientes B2B)',
         valor('Punto de equilibrio (clientes B2B)', lambda v: num(v, 0))],
        ['Clientes B2B en el año 5 (promedio del año)',
         num(filas['Clientes B2B (promedio del año, con contratos)'][-1], 1)],
        ['Necesidad máxima de financiamiento (USD)', num(necesidad)],
        ['Año de la necesidad máxima', str(anio)],
    ]


def tabla_comparada(escenarios):
    """Indicadores de todos los escenarios, uno por columna, en el orden de las hojas."""
    columnas = [filas_indicadores(filas) for filas in escenarios.values()]
    filas = [[fila[0][0], *(f[1] for f in fila)] for fila in zip(*columnas)]
    return tabla('comparada', None, 'Indicadores financieros comparados de los escenarios',
                 'l' + ' r' * len(escenarios), ['Indicador', *escenarios], filas)


def tabla_mercado(resumen, neutral):
    clientes = neutral['Clientes B2B (promedio del año, con contratos)'][-1]
    filas = [
        ['TAM', 'Todas las organizaciones de los cinco segmentos contratan su producto',
         num(resumen['TAM: todo el universo de organizaciones contrata'], 0)],
        ['SAM', 'Se excluyen los medios sin sitio web, que no pueden integrar la API',
         num(resumen['SAM: organizaciones que el producto y el canal atienden'], 0)],
        ['SOM', f'Ingresos del año 5 del escenario neutral, con {num(clientes, 1)} clientes en promedio',
         num(resumen['SOM: ingresos del año 5, escenario neutral'], 0)],
    ]
    return tabla('mercado', None, 'Mercado total, disponible y obtenible, en USD por año a precios del '
                 'escenario neutral. El TAM y el SAM suponen un año electoral',
                 'l p{8.4cm} r', ['Mercado', 'Criterio', 'USD por año'], filas)


def _bib():
    with open(os.path.join(ROOT, 'documento/biblio.bib')) as fh:
        return set(re.findall(r'@\w+\{([^,\s]+),', fh.read()))


def _latex(texto):
    return (str(texto).replace('\\', '').replace('&', '\\&').replace('%', '\\%')
            .replace('#', '\\#').replace('_', '\\_'))


FUENTES_INTERNAS = {
    'wiki/proyecto/recursos': 'registro de recursos del proyecto',
    'Cap. 3, tab:pricing': 'Tabla~\\ref{tab:pricing}',
    'Cap. 4, RNF-14': 'RNF-14, Sección~\\ref{sec:validacion}',
    'Decisión de #45': 'decisión del proyecto',
    'Cronograma del PFI': 'cronograma del trabajo',
    '[sin verificar]': '(a verificar)',
    '[sin verificar como compradores]': '(a verificar)',
}


def tabla_supuestos(supuestos):
    """Hoja Supuestos para el anexo: concepto, unidad, valor y fuente."""
    claves = _bib()

    def fuente(texto):
        t = _latex(texto or '—')
        for interna, publica in FUENTES_INTERNAS.items():
            t = t.replace(_latex(interna), publica)
        return re.sub(r'\b[A-Z][A-Za-z]+\d{4}[a-z]?\b',
                      lambda m: f'\\parencite{{{m[0]}}}' if m[0] in claves else m[0], t)

    def valor(v, unidad):
        if v is None:
            return '—'
        if isinstance(v, str):
            return _latex(v)
        if unidad == 'año':
            return str(int(v))
        if unidad == '%':
            return num(100 * v, 2 if 100 * v % 1 else 0) + '\\,\\%'
        return num(v, 0 if float(v).is_integer() else (4 if abs(v) < 1 else 2))

    filas = []
    for f in supuestos[2:]:
        if not any(f):
            continue
        if f[1] is None:  # título de grupo
            titulo = re.sub(r'scripts/\S+\.py', lambda m: '\\path{' + m[0].replace('\\_', '_') + '}', _latex(f[0]))
            filas.append(f'    \\multicolumn{{4}}{{@{{}}l}}{{\\textit{{{titulo}}}}} \\\\')
            continue
        unidad = f[2] or ''
        if f[3] is not None:
            v = valor(f[3], unidad)
        else:
            v = ' / '.join(valor(x, unidad) for x in f[4:7])
        # En una serie anual, la fuente va solo en la primera fila.
        origen = fuente(f[7]) if f[0] is not None else ''
        filas.append(f'    {_latex(f[1])} & {_latex(unidad)} & {v} & {origen} \\\\')
    cuerpo = '\n'.join(filas)
    return CABECERA + f"""{{\\footnotesize
\\begin{{longtable}}{{@{{}}>{{\\raggedright\\arraybackslash}}p{{5.2cm}} p{{2.1cm}} p{{2.4cm}} >{{\\raggedright\\arraybackslash}}p{{4.1cm}}@{{}}}}
  \\caption{{Supuestos del modelo financiero. Cuando el valor cambia por escenario se informa como optimista / neutral / pesimista. {FUENTE}}}
  \\label{{tab:financiero-supuestos}} \\\\
  \\hline
  \\textbf{{Supuesto}} & \\textbf{{Unidad}} & \\textbf{{Valor}} & \\textbf{{Fuente}} \\\\
  \\hline
  \\endfirsthead
  \\hline
  \\textbf{{Supuesto}} & \\textbf{{Unidad}} & \\textbf{{Valor}} & \\textbf{{Fuente}} \\\\
  \\hline
  \\endhead
  \\hline
  \\endfoot
{cuerpo}
\\end{{longtable}}
}}
"""


def en_texto(v, dec=0):
    """Como num, pero con el signo menos de la prosa en lugar de paréntesis."""
    return num(v, dec) if v >= 0 else '$-$' + num(-v, dec)


def cifras(resumen, supuestos, escenarios):
    """Una macro por cifra que el texto cita fuera de las tablas."""
    sup = {f[0]: f[3] for f in supuestos if f[0]}
    m = {
        'finTasaCapm': num(100 * resumen['Tasa CAPM de contraste']),
        'finHorasAnioCero': num(resumen['Valor de mercado de las horas de desarrollo del año 0, no imputado (USD)'], 0),
        'finTam': num(resumen['TAM: todo el universo de organizaciones contrata'], 0),
        'finSam': num(resumen['SAM: organizaciones que el producto y el canal atienden'], 0),
        'finSom': num(resumen['SOM: ingresos del año 5, escenario neutral'], 0),
        'finCargasSociales': num(100 * sup['cargas_sociales'], 0),
    }
    for nombre, filas in escenarios.items():
        p = 'fin' + nombre[:3].capitalize()
        necesidad, anio = necesidad_maxima(filas)
        tir_v = filas['TIR'][0]
        m |= {
            f'{p}Van': en_texto(filas['VAN (USD)'][0]),
            f'{p}VanResidual': en_texto(filas['VAN con valor residual (USD)'][0]),
            f'{p}Tir': tir_v if isinstance(tir_v, str) else en_texto(100 * tir_v, 2),
            f'{p}Inversion': num(filas['Costos totales'][0], 0),
            f'{p}Necesidad': num(necesidad, 0),
            f'{p}AnioNecesidad': str(anio),
        }
    return CABECERA + ''.join(f'\\newcommand{{\\{k}}}{{{v}}}\n' for k, v in m.items())


def economia_unitaria(filas_resumen):
    """Tabla de CAC y LTV del escenario neutral y macros de margen a cuota llena.

    Recalcula el LTV (margen × 12 / churn), el cociente LTV/CAC y el payback del CAC
    a partir de las columnas de la planilla; si no coinciden, devuelve errores.
    """
    por = {f[0]: f for f in filas_resumen if f[0]}
    cac = por['CAC (USD por alta de suscripción)'][1]
    errores, filas, ratios, paybacks = [], [], [], []
    for etiqueta, encabezado in SEGMENTOS:
        _, precio, costo, margen, churn, ltv, ratio, meses = por[f'Unitaria: {etiqueta}'][:8]
        for nombre, plan, propio in (('LTV', ltv, margen * 12 / churn), ('LTV/CAC', ratio, margen * 12 / churn / cac),
                                     ('payback del CAC', meses, cac / margen)):
            if not coincide(plan, propio):
                errores.append(f'ERROR: {nombre} de «{etiqueta}» no coincide: planilla = {plan}, recalculado = {propio}')
        ratios.append(ratio)
        paybacks.append(meses)
        filas.append([encabezado.replace('\\-', ''), num(precio, 0), num(costo, 0), num(margen, 0),
                      num(100 * churn, 0) + '\\,\\%', num(ltv, 0), num(ratio, 1), num(meses, 1)])
    tex = tabla('unitaria', None, 'Economía unitaria por segmento en el escenario neutral, en USD. El costo variable '
                'corresponde al consumo supuesto de consultas, dentro de la cuota; el LTV es el margen anual dividido '
                f'por el churn anual, y el CAC, de {num(cac, 0)} USD por alta, es común a todos los segmentos',
                'l r r r r r r r', ['Segmento', 'Precio', 'Costo var.', 'Margen', 'Churn', 'LTV', 'LTV/CAC',
                                    '\\textit{Payback} (meses)'], filas, letra='\\footnotesize\\setlength{\\tabcolsep}{3.5pt}')
    m = {'finCac': num(cac, 0), 'finLtvCacMin': num(min(ratios), 1), 'finLtvCacMax': num(max(ratios), 1),
         'finPaybackCacMax': num(max(paybacks), 1), 'finPaybackCacMin': num(min(paybacks), 1)}
    for nivel in ('intermedia', 'superior'):
        for i, esc in enumerate(('Opt', 'Neu', 'Pes')):
            m[f'finCosto{nivel.capitalize()}{esc}'] = num(por[f'API {nivel}: costo variable a cuota llena'][1 + i], 0)
            m[f'finMargen{nivel.capitalize()}{esc}'] = en_texto(por[f'API {nivel}: margen a cuota llena'][1 + i])
    return errores, tex, ''.join(f'\\newcommand{{\\{k}}}{{{v}}}\n' for k, v in m.items())


def tablas_del_escenario(nombre, filas):
    import unicodedata
    slug = unicodedata.normalize('NFKD', nombre).encode('ascii', 'ignore').decode().lower()
    esc = nombre.lower()
    anios = range(1, 6)

    altas = []
    for t in anios:
        fila = [str(t)]
        for etiqueta, _ in SEGMENTOS:
            a = entero(filas[f'{etiqueta}: altas'][t])
            fila.append(f'{"+" if a > 0 else ""}{a} ({entero(filas[f"{etiqueta}: activos al cierre"][t])})')
        fila.append(str(entero(filas['Organismos y ONG: contratos'][t])))
        altas.append(fila)

    flujo = []
    cero = [0] * 6
    for t in range(0, 6):
        capital = ((filas.get('Capital de trabajo: variación', cero)[t] or 0)
                   - (filas.get('Recupero del capital de trabajo', cero)[t] or 0))
        flujo.append([str(t)] + [num(v) for v in (
            filas['Ingresos totales'][t] or 0, filas['Costos totales'][t] or 0,
            filas.get('Impuesto a las ganancias', cero)[t] or 0, capital,
            filas['Flujo neto'][t] or 0, filas['Flujo descontado'][t] or 0,
            filas['Flujo neto acumulado'][t] or 0)])

    indicadores = filas_indicadores(filas)

    return slug, {
        'altas': tabla(slug, 'altas',
                       f'Altas de clientes B2B por año y segmento, escenario {esc}. Cada celda muestra '
                       'las altas del año y, entre paréntesis, los clientes activos al cierre; la '
                       'última columna corresponde a los contratos de monitoreo electoral',
                       'c *{6}{>{\\centering\\arraybackslash}p{1.8cm}}',
                       ['Año', *[e for _, e in SEGMENTOS], 'Contratos electorales'], altas),
        'flujo': tabla(slug, 'flujo', f'Flujo de caja del escenario {esc}, en USD. El flujo neto '
                       'resta de los ingresos los costos, el impuesto a las ganancias y la inversión en '
                       'capital de trabajo; en el año 5 esa inversión es negativa porque se recupera',
                       'c r r r r r r r',
                       ['Año', 'Ingresos', 'Costos', 'Impuesto', 'Cap. de trabajo', 'Flujo neto',
                        'Descontado', 'Acumulado'], flujo, letra='\\footnotesize\\setlength{\\tabcolsep}{3.5pt}'),
        'indicadores': tabla(slug, 'indicadores', f'Indicadores financieros del escenario {esc}',
                             'l r', ['Indicador', 'Valor'], indicadores),
    }


def main(args):
    parser = argparse.ArgumentParser(description='Valida el modelo financiero y exporta sus tablas LaTeX.')
    parser.add_argument('--autochequeo', action='store_true', help='solo corre el autochequeo')
    parser.add_argument('--planilla', default=PLANILLA, help='planilla recalculada a validar')
    parser.add_argument('--salida', default=SALIDA, help='directorio de las tablas .tex')
    opts = parser.parse_args(args)
    autochequeo()
    if opts.autochequeo:
        return 0
    planilla, salida = opts.planilla, opts.salida
    tasa, escenarios = leer(planilla)
    supuestos = leer_hoja(planilla, 'Supuestos')
    g = next((f[3] for f in supuestos if f[0] == 'crecimiento_perpetuo'), None)
    errores = []
    for nombre, filas in escenarios.items():
        for indicador, propio in recalcular(filas, tasa, g).items():
            if propio is None:
                propio = TEXTO_NULO[indicador]
            en_planilla = filas[indicador][0]
            if not coincide(en_planilla, propio):
                errores.append(f'ERROR: el indicador «{indicador}» no coincide en el escenario '
                               f'{nombre}: planilla = {en_planilla}, recalculado = {propio}')
    if errores:
        print('\n'.join(errores), file=sys.stderr)
        return 1
    os.makedirs(salida, exist_ok=True)
    for nombre, filas in escenarios.items():
        slug, tablas = tablas_del_escenario(nombre, filas)
        for clave, tex in tablas.items():
            ruta = os.path.join(salida, f'financiero-{slug}-{clave}.tex')
            with open(ruta, 'w') as fh:
                fh.write(tex)
            print(f'  escrito  {os.path.relpath(ruta, ROOT)}')
    extras = {'comparada': tabla_comparada(escenarios)}
    resumen = {f[0]: f[1] for f in leer_hoja(planilla, 'Resumen') if f[0]}
    if 'Neutral' in escenarios and 'TAM: todo el universo de organizaciones contrata' in resumen:
        extras['mercado'] = tabla_mercado(resumen, escenarios['Neutral'])
        extras['supuestos'] = tabla_supuestos(supuestos)
        extras['cifras'] = cifras(resumen, supuestos, escenarios)
        if 'CAC (USD por alta de suscripción)' in resumen:
            errores, extras['unitaria'], macros = economia_unitaria(leer_hoja(planilla, 'Resumen'))
            if errores:
                print('\n'.join(errores), file=sys.stderr)
                return 1
            extras['cifras'] += macros
    for clave, tex in extras.items():
        ruta = os.path.join(salida, f'financiero-{clave}.tex')
        with open(ruta, 'w') as fh:
            fh.write(tex)
        print(f'  escrito  {os.path.relpath(ruta, ROOT)}')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
