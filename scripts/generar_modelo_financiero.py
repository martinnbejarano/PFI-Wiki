#!/usr/bin/env python3
"""Genera la planilla del modelo financiero del PFI con fórmulas vivas.

Hoja «Supuestos»: el único lugar con números fijos. Cada supuesto lleva la
clave de wiki/negocio/analisis-financiero.md como nombre definido. Los
supuestos comunes van en la columna D; los que cambian por escenario, en las
columnas E (optimista), F (neutral) y G (pesimista).

Hojas «Optimista», «Neutral» y «Pesimista»: misma estructura, todo lo
derivado es fórmula. La celda B2 indica qué columna de escenario lee
(1 = optimista, 2 = neutral, 3 = pesimista); es lo único que las distingue.

Hoja «Resumen»: los indicadores de los tres escenarios, como fórmulas que
referencian las celdas de cada hoja.

openpyxl no calcula fórmulas: después de generar, hay que recalcular con
LibreOffice para que la planilla guarde los valores que lee el validador.

Uso:  .venv/bin/python scripts/generar_modelo_financiero.py
      soffice --headless --convert-to xlsx --outdir <tmp> wiki/negocio/modelo-financiero.xlsx
      (y copiar el resultado sobre el original)
"""
import os

import openpyxl
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANILLA = os.path.join(ROOT, 'wiki/negocio/modelo-financiero.xlsx')
WIKI = 'wiki/negocio/analisis-financiero.md'

# (clave, concepto, unidad, común | None, (opt, neu, pes) | None, fuente)
# Una serie anual es una lista de filas; la clave nombra el rango completo.
S = 'Supuesto propio'
SUPUESTOS = [
    ('# Horizonte, calendario y tasa',),
    ('horizonte', 'Horizonte de evaluación', 'años', 5, None, 'Decisión de #45'),
    ('anio_0', 'Año calendario del año 0', 'año', 2026, None, S),
    ('meses_anio', 'Meses por año', 'meses', 12, None, '—'),
    ('anio_electoral', [f'Año {t} con elección nacional (1 = sí)' for t in range(6)], '0/1',
     [0, 1, 0, 1, 0, 1], None, 'ConvencionConstituyente1994'),
    ('tasa_libre_riesgo', 'Bono del Tesoro de EE. UU. a 10 años', '%', 0.0517, None, 'FederalReserve2026'),
    ('riesgo_pais', 'Riesgo país de Argentina (EMBI, 628 pb)', '%', 0.0628, None, 'RavaBursatil2026'),
    ('prima_proyecto', 'Prima por proyecto nuevo', '%', 0.1355, None, 'Supuesto propio, Damodaran2010'),
    ('tasa_descuento', 'Tasa de descuento', '%',
     '=tasa_libre_riesgo+riesgo_pais+prima_proyecto', None, 'Suma de componentes'),
    ('erp_maduro', 'Prima de riesgo de un mercado maduro (contraste CAPM)', '%', 0.0423, None, 'Damodaran2026'),
    ('crp_argentina', 'Prima de riesgo país de Argentina (contraste CAPM)', '%', 0.0971, None, 'Damodaran2026'),
    ('tasa_capm', 'Tasa CAPM de contraste, con beta igual a 1', '%',
     '=tasa_libre_riesgo+erp_maduro+crp_argentina', None, 'Cálculo'),
    ('crecimiento_perpetuo', 'Crecimiento a perpetuidad para el valor residual', '%', 0.02, None,
     'Supuesto propio, no mayor al crecimiento de largo plazo'),
    ('# Recursos humanos',),
    ('sueldo_ars', 'Sueldo bruto mensual, developer semi senior', 'ARS/mes', 2501629, None, 'Sysarmy2026'),
    ('tc_encuesta', 'Tipo de cambio implícito de la encuesta', 'ARS/USD', 1397.88, None, 'Sysarmy2026'),
    ('sueldo_usd', 'Sueldo bruto mensual en USD', 'USD/mes', '=ROUND(sueldo_ars/tc_encuesta,0)', None, 'Cálculo'),
    ('sueldos_por_anio', 'Sueldos por año (12 + aguinaldo)', 'sueldos', 13, None, S),
    ('ajuste_sueldo', 'Ajuste anual del sueldo en USD', '%', 0.05, None, S),
    ('cargas_sociales', 'Cargas sociales del empleador (18 % + obra social 6 % + ART 2 %)', '%', 0.26, None,
     'CongresoNacion2019; ART: supuesto propio'),
    ('meses_anio_0', 'Meses de desarrollo del trabajo en el año 0 (no se imputan al flujo)', 'meses', 10, None,
     'Cronograma del PFI'),
    ('# Impuestos y capital de trabajo',),
    ('alicuota_ganancias', 'Impuesto a las ganancias, con quebrantos trasladables', '%', 0.35, None,
     'Alícuota máxima de la escala, CongresoNacion2021'),
    ('meses_cobro', 'Plazo de cobro de los clientes B2B y B2G', 'meses', 2, None, S),
    ('# Inversión del año 0 (además del desarrollo)',),
    ('inv_infraestructura_0', 'Infraestructura durante el PFI: API y clasificador en Railway (13 USD/mes × 12)', 'USD', 156, None,
     'Cap. 4, RNF-14; RailwayPrecios2026'),
    # ARS 8.500 (NIC Argentina, octubre de 2026) / ARS 1.528,76 por USD (BCRA2026, Com. A 3500).
    ('inv_dominio', 'Dominio .com.ar, primer año (ARS 8.500)', 'USD', 5.56, None, 'NICArgentina2026'),
    # 2 clases × 100 UMAPI × ARS 412,59 (octubre de 2026) = ARS 82.518 / 1.528,76.
    ('inv_marca_inpi', 'Solicitud de marca en el INPI, clases 9 y 42 (200 UMAPI)', 'USD', 54, None,
     'INPI2026, INPIUmapi2026'),
    ('inv_legal', 'Asesoría legal', 'USD', 1500, None, S),
    ('# Infraestructura y hosting (años 1 a 5)',),
    ('infra_base', 'Base comercial: Railway Pro, con 20 USD de uso incluidos (API y clasificador, 13 USD/mes)', 'USD/mes',
     20, None, 'RailwayPrecios2026'),
    ('infra_escalado', 'Uso adicional por cada bloque de usuarios activos (1 GB de memoria)', 'USD/mes', 10, None,
     'Supuesto propio; tarifa RailwayPrecios2026'),
    ('ua_escalado', 'Usuarios activos por bloque de escalado', 'UA', 10000, None, S),
    ('dominio_anual', 'Renovación del dominio .com.ar (ARS 8.500)', 'USD/año', 5.56, None, 'NICArgentina2026'),
    ('marketing', [f'Marketing, año {t}' for t in range(1, 6)], 'USD/año', None,
     [(6000, 4000, 2000), (9000, 6000, 3000), (12000, 8000, 4000),
      (15000, 10000, 5000), (18000, 12000, 6000)], S),
    ('# Costo variable por análisis',),
    ('precio_entrada', 'gpt-5.6-luna, entrada', 'USD/1 M tokens', 0.20, None, 'OpenAI2026'),
    ('precio_salida', 'gpt-5.6-luna, salida', 'USD/1 M tokens', 1.20, None, 'OpenAI2026'),
    ('precio_busqueda', 'Herramienta web_search', 'USD/llamada', 0.01, None, 'OpenAI2026'),
    ('tokens_por_millon', 'Tokens por millón', 'tokens', 1000000, None, '—'),
    ('tokens_entrada', 'Tokens de entrada por análisis', 'tokens', None, (6220, 11220, 18220), 'Supuesto propio; contraste: Cap. 4, RNF-14'),
    ('tokens_salida', 'Tokens de salida por análisis', 'tokens', None, (1250, 2000, 3600), 'Supuesto propio; contraste: Cap. 4, RNF-14'),
    ('busquedas', 'Llamadas a web_search por análisis', 'llamadas', None, (1, 2, 3), S),
    ('costo_analisis', 'Costo por análisis', 'USD', None, tuple(
        f'=({c}{{r}}*precio_entrada+{c}{{r_s}}*precio_salida)/tokens_por_millon+{c}{{r_b}}*precio_busqueda'
        for c in 'EFG'), 'Cálculo'),
    ('usuarios_activos', [f'Usuarios activos mensuales, año {t}' for t in range(1, 6)], 'UA', None,
     [(5000, 2000, 500), (20000, 8000, 2000), (50000, 20000, 5000),
      (100000, 40000, 8000), (150000, 60000, 10000)], S),
    ('analisis_por_ua_mes', 'Análisis por usuario activo por mes', 'análisis', None, (20, 12, 6), S),
    ('tope_mensual', 'Tope de análisis a demanda nuevos por instalación por mes', 'análisis', 5, None, 'Decisión de producto (RF-17)'),
    ('tasa_reuso', 'Análisis servidos desde la caché (las publicaciones virales concentran la demanda)', '%', None,
     (0.40, 0.30, 0.20), S),
    ('consultas_pro_mes', 'Demanda de consultas por cliente de la API intermedia por mes (se costea hasta la cuota)', 'consultas', None, (2000, 4000, 8000), S),
    ('consultas_ent_mes', 'Demanda de consultas por cliente de la API superior por mes (se costea hasta la cuota)', 'consultas', None, (20000, 40000, 80000), S),
    # El consumo que se costea es MIN(demanda, cuota): el excedente se cobra por consulta a un
    # precio mayor que su costo en los tres escenarios, de modo que omitirlo es conservador.
    ('cuota_intermedia', 'Cuota mensual incluida en la API intermedia', 'consultas', 5000, None, 'Cap. 3, tab:pricing'),
    ('cuota_superior', 'Cuota mensual incluida en la API superior', 'consultas', 40000, None, 'Cap. 3, tab:pricing'),
    ('precio_excedente', 'Precio por consulta excedente sobre la cuota (no se modela)', 'USD/consulta', 0.04, None,
     'Cap. 3, tab:pricing'),
    ('# Clientes B2B',),
    ('u_medios_total', 'Universo: socios activos de ADEPA, con y sin sitio web', 'organizaciones', 130, None, 'ADEPA2026'),
    ('u_medios_grandes', 'Universo: medios de gran porte', 'organizaciones', 10, None, 'ADEPA2026 (corte propio)'),
    ('u_medios_chicos', 'Universo: resto de medios de ADEPA con sitio web', 'organizaciones', 93, None, 'ADEPA2026'),
    ('u_verificadores', 'Universo: verificadores', 'organizaciones', 3, None, 'LatamChequea2026'),
    ('u_universidades', 'Universo: universidades', 'organizaciones', 149, None, 'SubsecretariaPoliticasUniversitarias2026'),
    ('u_organismos', 'Universo: organismos electorales y ONG', 'organizaciones', 32, None, 'MinisterioInterior2026'),
    ('u_agencias', 'Universo: agencias', 'organizaciones', 120, None, 'AgenciasArgentinas2026'),
    ('precio_medios_grandes', 'Precio: API nivel superior, medios de gran porte', 'USD/mes', None, (3000, 2500, 1500), 'Cap. 3, tab:pricing'),
    ('precio_medios_chicos', 'Precio: API nivel intermedio', 'USD/mes', 200, None, 'Cap. 3, tab:pricing'),
    ('precio_verificadores', 'Precio: panel de tendencias + API intermedia', 'USD/mes', 500, None, 'Cap. 3, tab:pricing'),
    ('precio_universidades', 'Precio: acceso histórico', 'USD/mes', 100, None, 'Cap. 3, tab:pricing'),
    ('precio_organismos', 'Precio: reporte de monitoreo electoral', 'USD/contrato', None, (30000, 20000, 10000), 'Cap. 3, tab:pricing'),
    ('precio_agencias', 'Precio: API nivel superior, agencias', 'USD/mes', None, (3000, 2500, 1500), 'Cap. 3, tab:pricing'),
    # Criterio único para los tres escenarios: las ventas empiezan el primer año con este piso de UA.
    ('ua_inicio_b2b', 'Usuarios activos mensuales desde los que empiezan las ventas B2B', 'UA', 5000, None, S),
    ('capt_medios_grandes', 'Captación anual: medios de gran porte', '%', None, (0.20, 0.10, 0.05), S),
    ('capt_medios_chicos', 'Captación anual: resto de medios de ADEPA', '%', None, (0.10, 0.05, 0.02), S),
    ('capt_verificadores', 'Captación anual: verificadores', '%', None, (0.40, 0.30, 0.20), S),
    ('capt_universidades', 'Captación anual: universidades', '%', None, (0.08, 0.04, 0.02), S),
    ('capt_organismos', 'Fracción de organismos que contrata en año electoral', '%', None, (0.10, 0.06, 0.03), S),
    ('capt_agencias', 'Captación anual: agencias', '%', None, (0.05, 0.02, 0.01), S),
    ('churn_medios_grandes', 'Churn anual: medios de gran porte', '%', None, (0.10, 0.15, 0.25), S),
    ('churn_medios_chicos', 'Churn anual: resto de medios de ADEPA', '%', None, (0.15, 0.20, 0.30), S),
    ('churn_verificadores', 'Churn anual: verificadores', '%', None, (0.00, 0.10, 0.20), S),
    ('churn_universidades', 'Churn anual: universidades', '%', None, (0.05, 0.10, 0.15), S),
    ('churn_agencias', 'Churn anual: agencias', '%', None, (0.15, 0.20, 0.30), S),
    ('# Sensibilidad (1 = caso base)',),
    ('sens_costo_analisis', 'Multiplicador del costo por análisis', 'veces', 1, None, '—'),
    ('sens_precios', 'Multiplicador de los precios B2B', 'veces', 1, None, '—'),
    ('sens_captacion', 'Multiplicador de las tasas de captación', 'veces', 1, None, '—'),
]

# (sufijo de la clave, etiqueta en la hoja de escenario, producto)
SEGMENTOS = [
    ('medios_grandes', 'Medios de gran porte', 'API nivel superior'),
    ('medios_chicos', 'Medios socios de ADEPA (resto)', 'API nivel intermedio'),
    ('verificadores', 'Verificadores', 'panel de tendencias y API intermedia'),
    ('universidades', 'Universidades y observatorios', 'acceso histórico'),
    ('agencias', 'Agencias', 'API nivel superior'),
]

NEGRITA = Font(bold=True)
COLS = 'BCDEFG'  # años 0 a 5


def hoja_supuestos(wb):
    ws = wb.active
    ws.title = 'Supuestos'
    ws.append(['clave', 'concepto', 'unidad', 'común', 'optimista', 'neutral', 'pesimista', 'fuente'])
    for c in ws[1]:
        c.font = NEGRITA
    ws.append([f'Fuente de cada supuesto: {WIKI}. Clave bibliográfica = documento/biblio.bib.'])
    filas = {}
    for s in SUPUESTOS:
        if len(s) == 1:
            ws.append([])
            ws.append([s[0][2:]])
            ws.cell(ws.max_row, 1).font = NEGRITA
            continue
        clave, concepto, unidad, comun, esc, fuente = s
        inicio = ws.max_row + 1
        conceptos = concepto if isinstance(concepto, list) else [concepto]
        for i, texto in enumerate(conceptos):
            c = comun[i] if isinstance(comun, list) else comun
            e = esc[i] if isinstance(esc, list) else esc
            ws.append([clave if i == 0 else None, texto, unidad, c, *(e or (None,) * 3), fuente])
        fin = ws.max_row
        filas[clave] = inicio
        cols = ('$D', '$D') if comun is not None else ('$E', '$G')
        ref = f"Supuestos!{cols[0]}${inicio}:{cols[1]}${fin}"
        if comun is not None and inicio == fin:
            ref = f"Supuestos!$D${inicio}"
        wb.defined_names[clave] = DefinedName(clave, attr_text=ref)
    # costo_analisis referencia las filas de sus insumos por escenario.
    r = filas['costo_analisis']
    for c in 'EFG':
        ws[f'{c}{r}'] = ws[f'{c}{r}'].value.format(
            r=filas['tokens_entrada'], r_s=filas['tokens_salida'], r_b=filas['busquedas'])
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 62
    ws.column_dimensions['H'].width = 40


def hoja_escenario(wb, nombre, columna):
    ws = wb.create_sheet(nombre)
    ws['A1'] = f'Escenario {nombre.lower()} (USD nominales; todo lo derivado es fórmula)'
    ws['A1'].font = NEGRITA
    ws['A2'] = 'Columna de escenario en Supuestos (1 = optimista, 2 = neutral, 3 = pesimista)'
    ws['B2'] = columna
    ws.column_dimensions['A'].width = 60
    fila = {}

    def poner(etiqueta, formula, anios=range(6)):
        """formula(c, p, t): columna del año, columna del año anterior, celda del año."""
        r = fila['_r'] = ws.max_row + 1
        ws.cell(r, 1, etiqueta)
        for i in anios:
            c = COLS[i]
            p = COLS[i - 1] if i else None
            ws[f'{c}{r}'] = '=' + formula(c, p, f'{c}${fila.get("Año", r)}')
        fila[etiqueta] = r
        return r

    def titulo(texto):
        ws.append([])
        ws.append([texto])
        ws.cell(ws.max_row, 1).font = NEGRITA

    esc = '$B$2'
    X = lambda clave: f'INDEX({clave},1,{esc})'      # supuesto por escenario
    XT = lambda clave, t: f'INDEX({clave},{t},{esc})'  # serie anual por escenario, años 1 a 5
    R = lambda etiqueta, c: f'{c}{fila[etiqueta]}'
    OPER = range(1, 6)

    ws.append([])
    poner('Año', lambda c, p, t: 'COLUMN()-2')
    poner('Año calendario', lambda c, p, t: f'anio_0+{t}')
    poner('Año electoral (1 = sí)', lambda c, p, t: f'INDEX(anio_electoral,{t}+1)')
    poner('Usuarios activos mensuales', lambda c, p, t: XT('usuarios_activos', t), OPER)
    poner('Ventas B2B habilitadas (1 = sí)', lambda c, p, t: (
        f'IF({R("Usuarios activos mensuales", c)}>=ua_inicio_b2b,1,0)'), OPER)

    titulo('Clientes B2B')
    for suf, et, _ in SEGMENTOS:
        ini, fin = f'{et}: activos al inicio', f'{et}: activos al cierre'
        cierre = ws.max_row + 4  # fila de «activos al cierre», tres filas más abajo
        poner(ini, lambda c, p, t: f'N({p}{cierre})', OPER)
        poner(f'{et}: bajas', lambda c, p, t: f'ROUND({X("churn_" + suf)}*{R(ini, c)},0)', OPER)
        poner(f'{et}: altas', lambda c, p, t: (
            f'{R("Ventas B2B habilitadas (1 = sí)", c)}*ROUND({X("capt_" + suf)}*sens_captacion*(u_{suf}-{R(ini, c)}),0)'), OPER)
        poner(fin, lambda c, p, t: f'{R(ini, c)}+{R(et + ": altas", c)}-{R(et + ": bajas", c)}', OPER)
        poner(f'{et}: promedio del año', lambda c, p, t: f'({R(ini, c)}+{R(fin, c)})/2', OPER)
    poner('Organismos y ONG: contratos', lambda c, p, t: (
        f'{R("Año electoral (1 = sí)", c)}*{R("Ventas B2B habilitadas (1 = sí)", c)}'
        f'*ROUND({X("capt_organismos")}*sens_captacion*u_organismos,0)'), OPER)
    poner('Clientes B2B (promedio del año, con contratos)', lambda c, p, t: '+'.join(
        [R(f'{et}: promedio del año', c) for _, et, _ in SEGMENTOS] + [R('Organismos y ONG: contratos', c)]), OPER)

    titulo('Ingresos por producto')
    ingresos = []
    for suf, et, prod in SEGMENTOS:
        precio = X(f'precio_{suf}') if suf in ('medios_grandes', 'agencias') else f'precio_{suf}'
        ingresos.append(poner(f'Ingresos: {prod} ({et})', lambda c, p, t: (
            f'{R(et + ": promedio del año", c)}*{precio}*sens_precios*meses_anio'), OPER))
    ingresos.append(poner('Ingresos: reportes de monitoreo electoral, organismos y ONG', lambda c, p, t: (
        f'{R("Organismos y ONG: contratos", c)}*{X("precio_organismos")}*sens_precios'), OPER))
    poner('Ingresos totales', lambda c, p, t: f'SUM({c}{ingresos[0]}:{c}{ingresos[-1]})')

    titulo('Costos por rubro')
    ua = lambda c: R('Usuarios activos mensuales', c)
    prom = lambda et, c: R(f'{et}: promedio del año', c)
    variable = f'meses_anio*(1-{X("tasa_reuso")})*{X("costo_analisis")}*sens_costo_analisis'
    costos = [
        # Las horas de desarrollo del año 0 no se imputan (criterio de la cátedra):
        # su valor de mercado queda en la hoja Resumen, para la defensa.
        poner('Costo: recursos humanos (desarrollador semi senior, con cargas sociales)', lambda c, p, t: (
            f'sueldo_usd*sueldos_por_anio*(1+cargas_sociales)*(1+ajuste_sueldo)^{t}'), OPER),
        poner('Costo: infraestructura', lambda c, p, t: (
            'inv_infraestructura_0' if c == 'B'
            else f'meses_anio*(infra_base+infra_escalado*{ua(c)}/ua_escalado)+dominio_anual')),
        poner('Costo: marketing', lambda c, p, t: XT('marketing', t), OPER),
        poner('Costo: dominio, marca y asesoría legal',
              lambda c, p, t: 'inv_dominio+inv_marca_inpi+inv_legal', [0]),
        poner('Costo variable: análisis de usuarios gratuitos', lambda c, p, t: (
            f'{ua(c)}*MIN({X("analisis_por_ua_mes")},tope_mensual)*{variable}'), OPER),
        poner('Costo variable: consultas de la API B2B', lambda c, p, t: (
            f'(({prom("Medios socios de ADEPA (resto)", c)}+{prom("Verificadores", c)})*MIN({X("consultas_pro_mes")},cuota_intermedia)'
            f'+({prom("Medios de gran porte", c)}+{prom("Agencias", c)})*MIN({X("consultas_ent_mes")},cuota_superior))*{variable}'), OPER),
    ]
    poner('Costos totales', lambda c, p, t: f'SUM({c}{costos[0]}:{c}{costos[-1]})')

    titulo('Impuesto y capital de trabajo')
    poner('Resultado antes de impuestos', lambda c, p, t: f'{R("Ingresos totales", c)}-{R("Costos totales", c)}')
    rai = lambda c: R('Resultado antes de impuestos', c)
    # Quebranto acumulado (≤ 0): la pérdida de un año se descuenta de las ganancias siguientes.
    poner('Quebranto acumulado', lambda c, p, t: (
        f'MIN(0,{rai(c)})' if c == 'B' else f'MIN(0,{p}{fila["_r"]}+{rai(c)})'))
    poner('Impuesto a las ganancias', lambda c, p, t: (
        f'alicuota_ganancias*MAX(0,{rai(c)})' if c == 'B'
        else f'alicuota_ganancias*MAX(0,{p}{fila["Quebranto acumulado"]}+{rai(c)})'))
    poner('Capital de trabajo', lambda c, p, t: f'meses_cobro/meses_anio*{R("Ingresos totales", c)}')
    poner('Capital de trabajo: variación', lambda c, p, t: (
        R('Capital de trabajo', c) if c == 'B' else f'{R("Capital de trabajo", c)}-{p}{fila["Capital de trabajo"]}'))
    # Las cuentas a cobrar se liquidan al cierre del horizonte (Sapag Chain et al., 2014).
    poner('Recupero del capital de trabajo', lambda c, p, t: R('Capital de trabajo', c), [5])

    titulo('Flujo de caja')
    poner('Flujo neto', lambda c, p, t: (
        f'{rai(c)}-{R("Impuesto a las ganancias", c)}-{R("Capital de trabajo: variación", c)}'
        f'+N({R("Recupero del capital de trabajo", c)})'))
    poner('Factor de descuento', lambda c, p, t: f'1/(1+tasa_descuento)^{t}')
    poner('Flujo descontado', lambda c, p, t: f'{R("Flujo neto", c)}*{R("Factor de descuento", c)}')
    poner('Flujo neto acumulado', lambda c, p, t: (
        R('Flujo neto', c) if c == 'B' else f'{p}{fila["_r"]}+{R("Flujo neto", c)}'))
    poner('Flujo descontado acumulado', lambda c, p, t: (
        R('Flujo descontado', c) if c == 'B' else f'{p}{fila["_r"]}+{R("Flujo descontado", c)}'))
    for tipo, acum, flujo in (('simple', 'Flujo neto acumulado', 'Flujo neto'),
                              ('descontado', 'Flujo descontado acumulado', 'Flujo descontado')):
        poner(f'Recupero {tipo}: a + b/c en el año del cruce', lambda c, p, t: (
            f'IF(AND({p}{fila[acum]}<0,{R(acum, c)}>=0),{t}-1+(-{p}{fila[acum]})/{R(flujo, c)},"")'), OPER)

    titulo('Indicadores')
    rango = lambda et: f'B{fila[et]}:G{fila[et]}'
    ind = [
        ('VAN (USD)', f'SUM({rango("Flujo descontado")})'),
        ('TIR', f'IFERROR(IRR({rango("Flujo neto")}),"no definida")'),
        ('Payback simple (años)', 'IF(COUNT({0})=0,"no se recupera",MIN({0}))'.format(
            rango('Recupero simple: a + b/c en el año del cruce').replace('B', 'C', 1))),
        ('Payback descontado (años)', 'IF(COUNT({0})=0,"no se recupera",MIN({0}))'.format(
            rango('Recupero descontado: a + b/c en el año del cruce').replace('B', 'C', 1))),
        # Gordon sobre el flujo operativo del año 5 (sin el recupero del capital de
        # trabajo). Con flujo operativo negativo se liquida: valor residual nulo.
        ('VAN con valor residual (USD)', (
            f'SUM({rango("Flujo descontado")})+MAX(0,G{fila["Flujo neto"]}-G{fila["Recupero del capital de trabajo"]})'
            f'*(1+crecimiento_perpetuo)/(tasa_descuento-crecimiento_perpetuo)/(1+tasa_descuento)^5')),
        ('Ingreso medio por cliente B2B en el año 5 (USD)', (
            f'IFERROR(G{fila["Ingresos totales"]}/G{fila["Clientes B2B (promedio del año, con contratos)"]},"no definido")')),
    ]
    for etiqueta, f in ind:
        ws.append([etiqueta, '=' + f])
        fila[etiqueta] = ws.max_row
    ws.append(['Punto de equilibrio (clientes B2B)',
               f'=IFERROR(ROUNDUP(G{fila["Costos totales"]}/B{ws.max_row},0),"no definido")'])
    fila['Punto de equilibrio (clientes B2B)'] = ws.max_row
    ws[f'B{fila["TIR"]}'].number_format = '0.00%'
    return fila


ESCENARIOS = [('Optimista', 1), ('Neutral', 2), ('Pesimista', 3)]
INDICADORES = ['VAN (USD)', 'TIR', 'Payback simple (años)', 'Payback descontado (años)',
               'Punto de equilibrio (clientes B2B)',
               'Ingreso medio por cliente B2B en el año 5 (USD)', 'VAN con valor residual (USD)']


def hoja_resumen(wb, filas):
    """filas: la fila de cada indicador, igual en las tres hojas de escenario.

    Debajo de los indicadores van el mercado (TAM, SAM y SOM, a precios del
    escenario neutral) y dos cifras de contraste para la defensa.
    """
    ws = wb.create_sheet('Resumen', 1)
    ws.append(['Indicador', *(n for n, _ in ESCENARIOS)])
    for c in ws[1]:
        c.font = NEGRITA
    for etiqueta in INDICADORES:
        ws.append([etiqueta, *(f'={n}!B{filas[etiqueta]}' for n, _ in ESCENARIOS)])
    for col in 'BCD':
        ws[f'{col}3'].number_format = '0.00%'
    neu = lambda clave: f'INDEX({clave},1,2)'
    anual = lambda medios_resto: (
        f'=meses_anio*(u_medios_grandes*{neu("precio_medios_grandes")}+{medios_resto}*precio_medios_chicos'
        f'+u_verificadores*precio_verificadores+u_universidades*precio_universidades'
        f'+u_agencias*{neu("precio_agencias")})+u_organismos*{neu("precio_organismos")}')
    ws.append([])
    ws.append(['Mercado, a precios del escenario neutral (USD por año)'])
    ws.cell(ws.max_row, 1).font = NEGRITA
    ws.append(['TAM: todo el universo de organizaciones contrata', anual('(u_medios_total-u_medios_grandes)')])
    ws.append(['SAM: organizaciones que el producto y el canal atienden', anual('u_medios_chicos')])
    ws.append(['SOM: ingresos del año 5, escenario neutral', f'=Neutral!G{filas["Ingresos totales"]}'])
    ws.append([])
    ws.append(['Contraste para la defensa'])
    ws.cell(ws.max_row, 1).font = NEGRITA
    ws.append(['Valor de mercado de las horas de desarrollo del año 0, no imputado (USD)',
               '=meses_anio_0*sueldo_usd*(1+cargas_sociales)*sueldos_por_anio/meses_anio'])
    ws.append(['Tasa CAPM de contraste', '=tasa_capm'])

    # Margen por cliente con la cuota consumida por completo, en los tres escenarios.
    ws.append([])
    ws.append(['Margen mensual por cliente a cuota llena (USD)', *(n for n, _ in ESCENARIOS)])
    ws.cell(ws.max_row, 1).font = NEGRITA
    costo_q = lambda k: f'(1-INDEX(tasa_reuso,1,{k}))*INDEX(costo_analisis,1,{k})*sens_costo_analisis'
    for nivel, cuota, precio in (('intermedia', 'cuota_intermedia', lambda k: 'precio_medios_chicos'),
                                 ('superior', 'cuota_superior', lambda k: f'INDEX(precio_medios_grandes,1,{k})')):
        ws.append([f'API {nivel}: costo variable a cuota llena', *(f'={cuota}*{costo_q(k)}' for _, k in ESCENARIOS)])
        ws.append([f'API {nivel}: margen a cuota llena',
                   *(f'={precio(k)}*sens_precios-{cuota}*{costo_q(k)}' for _, k in ESCENARIOS)])

    # CAC y LTV del escenario neutral. El CAC imputa todo el marketing de los años 1 a 5 a las
    # altas de suscripciones (cota superior: el marketing también capta usuarios ciudadanos).
    ws.append([])
    ws.append(['Economía unitaria, escenario neutral (USD)', 'Precio mensual', 'Costo variable mensual',
               'Margen mensual', 'Churn anual', 'LTV', 'LTV/CAC', 'Payback del CAC (meses)'])
    ws.cell(ws.max_row, 1).font = NEGRITA
    altas = '+'.join(f'SUM(Neutral!C{filas[et + ": altas"]}:G{filas[et + ": altas"]})' for _, et, _ in SEGMENTOS)
    ws.append(['CAC (USD por alta de suscripción)', f'=SUM(Neutral!C{filas["Costo: marketing"]}:G{filas["Costo: marketing"]})/({altas})'])
    cac = f'$B${ws.max_row}'
    consumo = {'medios_grandes': 'MIN(INDEX(consultas_ent_mes,1,2),cuota_superior)',
               'agencias': 'MIN(INDEX(consultas_ent_mes,1,2),cuota_superior)',
               'medios_chicos': 'MIN(INDEX(consultas_pro_mes,1,2),cuota_intermedia)',
               'verificadores': 'MIN(INDEX(consultas_pro_mes,1,2),cuota_intermedia)',
               'universidades': '0'}
    for suf, et, _ in SEGMENTOS:
        r = ws.max_row + 1
        precio = f'INDEX(precio_{suf},1,2)' if suf in ('medios_grandes', 'agencias') else f'precio_{suf}'
        ws.append([f'Unitaria: {et}', f'={precio}*sens_precios', f'={consumo[suf]}*{costo_q(2)}',
                   f'=B{r}-C{r}', f'=INDEX(churn_{suf},1,2)', f'=D{r}*meses_anio/E{r}',
                   f'=F{r}/{cac}', f'={cac}/D{r}'])
    ws.column_dimensions['A'].width = 50


def main():
    wb = openpyxl.Workbook()
    hoja_supuestos(wb)
    for nombre, columna in ESCENARIOS:
        filas = hoja_escenario(wb, nombre, columna)
    hoja_resumen(wb, filas)
    wb.save(PLANILLA)
    print(f'  escrito  {os.path.relpath(PLANILLA, ROOT)} (sin valores: recalcular con LibreOffice)')


if __name__ == '__main__':
    main()
