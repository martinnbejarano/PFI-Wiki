"""Tests de scripts/validar_modelo_financiero.py a través de su interfaz de línea de comandos.

Uso:  .venv/bin/python -m unittest scripts/test_validar_modelo_financiero.py
"""
import os, subprocess, sys, tempfile, unittest

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPT = os.path.join(ROOT, 'scripts/validar_modelo_financiero.py')
PLANILLA = os.path.join(ROOT, 'wiki/negocio/modelo-financiero.xlsx')
SEGMENTOS = ['Medios de gran porte', 'Medios socios de ADEPA (resto)', 'Verificadores',
             'Universidades y observatorios', 'Agencias']


def correr(*args):
    return subprocess.run([sys.executable, SCRIPT, *args], capture_output=True, text=True)


def hoja_de_juguete(wb, nombre, neto, descontado, anio5, indicadores):
    """Hoja de escenario con valores (sin fórmulas). anio5 = (ingresos, costos, clientes)."""
    esc = wb.create_sheet(nombre)
    acumulado = [sum(neto[:t + 1]) for t in range(6)]
    filas = [
        ('Año', [0, 1, 2, 3, 4, 5]),
        ('Ingresos totales', [0, 60, 60, 0, 0, anio5[0]]),
        ('Costos totales', [100, 0, 0, 0, 0, anio5[1]]),
        ('Flujo neto', neto),
        ('Flujo descontado', descontado),
        ('Flujo neto acumulado', acumulado),
        ('Clientes B2B (promedio del año, con contratos)', [0, 0, 0, 0, 0, anio5[2]]),
        ('Organismos y ONG: contratos', [None, 0, 0, 0, 0, 2]),
    ]
    for segmento in SEGMENTOS:
        filas.append((f'{segmento}: altas', [None, 1, 0, 0, 0, 0]))
        filas.append((f'{segmento}: activos al cierre', [None, 1, 1, 1, 1, 1]))
    for etiqueta, valores in filas:
        esc.append([etiqueta, *valores])
    esc.append([])
    for etiqueta, valor in zip(INDICADORES, indicadores):
        esc.append([etiqueta, valor])


INDICADORES = ['VAN (USD)', 'TIR', 'Payback simple (años)', 'Payback descontado (años)',
               'Punto de equilibrio (clientes B2B)']


def planilla_de_juguete(ruta, van_planilla, tres_escenarios=False):
    """Planilla con valores (sin fórmulas) y flujos de resultado conocido, al 25 %.

    Neutral: -100, 60, 60, 0, 0, 0 => VAN = -100 + 48 + 38,4 = -13,6;
    acumulado -100, -40, 20 => payback = 1 + 40/60; el descontado no se recupera.
    TIR = 13,0662 % (raíz de 60x² + 60x - 100 = 0, x = 1/(1+TIR)).
    Año 5: costos 1.000, ingresos 1.000 y 4 clientes => 250 por cliente => 4 clientes.

    Optimista: -100, 150 => VAN = -100 + 120 = 20; TIR = 50 %; payback = 100/150;
    descontado = 100/120. Año 5: costos 1.000, ingresos 2.000, 4 clientes => 2.

    Pesimista: -100, -10, -5 => VAN = -100 - 8 - 3,2 = -111,2; sin cambio de signo
    no hay TIR ni recupero. Año 5 sin ingresos ni clientes => sin punto de equilibrio.
    """
    wb = openpyxl.Workbook()
    sup = wb.active
    sup.title = 'Supuestos'
    sup.append(['clave', 'concepto', 'unidad', 'común', 'Opt', 'Neu', 'Pes'])
    sup.append(['tasa_descuento', 'Tasa de descuento', '%', 0.25])
    if tres_escenarios:
        hoja_de_juguete(wb, 'Optimista', [-100, 150, 0, 0, 0, 0], [-100, 120, 0, 0, 0, 0],
                        (2000, 1000, 4), [20, 0.5, 100 / 150, 100 / 120, 2])
    hoja_de_juguete(wb, 'Neutral', [-100, 60, 60, 0, 0, 0], [-100, 48, 38.4, 0, 0, 0],
                    (1000, 1000, 4), [van_planilla, 0.130662, 1 + 40 / 60, 'no se recupera', 4])
    if tres_escenarios:
        hoja_de_juguete(wb, 'Pesimista', [-100, -10, -5, 0, 0, 0], [-100, -8, -3.2, 0, 0, 0],
                        (0, 1000, 0), [-111.2, 'no definida', 'no se recupera', 'no se recupera',
                                       'no definido'])
    wb.save(ruta)


class ValidarModeloFinanciero(unittest.TestCase):

    def test_autochequeo_con_flujos_de_juguete_pasa(self):
        r = correr('--autochequeo')
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn('autochequeo: ok', r.stdout)

    def test_argumentos_invalidos_no_validan_ni_escriben_tablas(self):
        with tempfile.TemporaryDirectory() as d:
            for args in (['--help'], ['--desconocido'], ['--planilla']):
                r = correr(*args, '--salida', d)
                self.assertNotIn('escrito', r.stdout, args)
                self.assertNotIn('Traceback', r.stderr, args)
            self.assertEqual(os.listdir(d), [])
            self.assertEqual(correr('--help').returncode, 0)
            self.assertEqual(correr('--desconocido').returncode, 2)

    def test_indicador_que_no_coincide_termina_con_error_y_lo_nombra(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, 'm.xlsx')
            planilla_de_juguete(ruta, van_planilla=-13.0)  # el correcto es -13,6
            r = correr('--planilla', ruta, '--salida', d)
            self.assertEqual(r.returncode, 1)
            self.assertIn('VAN', r.stderr)
            self.assertIn('Neutral', r.stderr)
            self.assertFalse(os.path.exists(os.path.join(d, 'financiero-neutral-flujo.tex')))

    def test_planilla_consistente_genera_las_tablas_del_escenario(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, 'm.xlsx')
            planilla_de_juguete(ruta, van_planilla=-13.6)
            r = correr('--planilla', ruta, '--salida', d)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            tablas = {}
            for nombre in ('altas', 'flujo', 'indicadores'):
                with open(os.path.join(d, f'financiero-neutral-{nombre}.tex')) as fh:
                    tablas[nombre] = fh.read()
            for tex in tablas.values():
                self.assertIn(r'\begin{table}[t]', tex)
                self.assertLess(tex.index(r'\caption'), tex.index(r'\begin{tabular}'))
                self.assertIn('Fuente: elaboración propia', tex)
            flujo, ind = tablas['flujo'], tablas['indicadores']
            self.assertIn('1.000,00', flujo)        # punto de miles y coma decimal
            self.assertIn('(100,00)', flujo)        # negativos entre paréntesis
            self.assertIn('(13,60)', ind)           # VAN
            self.assertIn('13,07', ind)             # TIR en %
            self.assertIn('1,67', ind)              # payback a mitad de año
            self.assertIn('no se recupera', ind)    # payback descontado
            self.assertIn('+1 (1)', tablas['altas'])

    def test_tres_escenarios_generan_sus_tablas_y_la_comparada(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, 'm.xlsx')
            planilla_de_juguete(ruta, van_planilla=-13.6, tres_escenarios=True)
            r = correr('--planilla', ruta, '--salida', d)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            for esc in ('optimista', 'neutral', 'pesimista'):
                for nombre in ('altas', 'flujo', 'indicadores'):
                    self.assertTrue(os.path.exists(os.path.join(d, f'financiero-{esc}-{nombre}.tex')))
            with open(os.path.join(d, 'financiero-comparada.tex')) as fh:
                tex = fh.read()
            self.assertIn(r'\begin{table}[t]', tex)
            self.assertLess(tex.index(r'\caption'), tex.index(r'\begin{tabular}'))
            self.assertIn('Fuente: elaboración propia', tex)
            self.assertIn(r'\label{tab:financiero-comparada}', tex)
            # columnas en orden optimista, neutral, pesimista
            self.assertLess(tex.index('Optimista'), tex.index('Neutral'))
            self.assertLess(tex.index('Neutral'), tex.index('Pesimista'))
            van = next(l for l in tex.splitlines() if l.strip().startswith('VAN'))
            self.assertEqual([c.strip(' \\') for c in van.split('&')[1:]], ['20,00', '(13,60)', '(111,20)'])
            tir = next(l for l in tex.splitlines() if l.strip().startswith('TIR'))
            self.assertEqual([c.strip(' \\') for c in tir.split('&')[1:]], ['50,00', '13,07', 'no definida'])
            self.assertIn('0,67', tex)              # payback optimista
            self.assertIn('no definido', tex)       # punto de equilibrio pesimista

    def test_la_planilla_versionada_valida_contra_el_calculo_independiente(self):
        with tempfile.TemporaryDirectory() as d:
            r = correr('--planilla', PLANILLA, '--salida', d)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            for esc in ('optimista', 'neutral', 'pesimista'):
                self.assertTrue(os.path.exists(os.path.join(d, f'financiero-{esc}-indicadores.tex')))
            self.assertTrue(os.path.exists(os.path.join(d, 'financiero-comparada.tex')))

    def test_la_hoja_resumen_referencia_los_indicadores_de_los_tres_escenarios(self):
        formulas = openpyxl.load_workbook(PLANILLA)['Resumen']
        valores = openpyxl.load_workbook(PLANILLA, data_only=True)
        resumen = {f[0]: f[1:4] for f in valores['Resumen'].iter_rows(values_only=True) if f[0]}
        self.assertEqual(list(resumen['Indicador']), ['Optimista', 'Neutral', 'Pesimista'])
        for hoja, col in (('Optimista', 'B'), ('Neutral', 'C'), ('Pesimista', 'D')):
            propios = {f[0]: f[1] for f in valores[hoja].iter_rows(values_only=True) if f[0]}
            for i, etiqueta in enumerate(INDICADORES):
                self.assertEqual(resumen[etiqueta][' BCD'.index(col) - 1], propios[etiqueta])
                self.assertTrue(str(formulas[f'{col}{i + 2}'].value).startswith(f'={hoja}!'))


if __name__ == '__main__':
    unittest.main()
